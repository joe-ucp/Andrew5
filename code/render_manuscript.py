"""Generate matching Word and LaTeX sources from authoritative scientific content."""
from pathlib import Path
import json,re
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];M=ROOT/'manuscript'

def esc(s):
    s=str(s)
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(mapping.get(c,c) for c in s)

def no_split(row):
    pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))

def main():
    data=json.loads((M/'paper_content.json').read_text(encoding='utf8'));doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=sec.right_margin=Inches(.85)
    style=doc.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(11)
    style.paragraph_format.line_spacing=1.08;style.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',18),('Heading 1',13),('Heading 2',11)]:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6)
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.core_properties.author=data['author'];doc.core_properties.last_modified_by=data['author'];doc.core_properties.title=data['title'];doc.core_properties.subject='Measurement and transferability of a human islet reprogramming signature'
    p=doc.add_paragraph(data['title'],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(data['author']+'\n'+data['date']);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    tex=[r'\documentclass[11pt,letterpaper]{article}',r'\usepackage[margin=0.85in]{geometry}',r'\usepackage[T1]{fontenc}',r'\usepackage[utf8]{inputenc}',r'\usepackage{lmodern,amsmath,amssymb,graphicx,tabularx,array,booktabs,float,hyperref}',r'\hypersetup{colorlinks=true,urlcolor=blue,pdfauthor={UCP Technology LLC},pdftitle={'+esc(data['title'])+'}}',r'\setlength{\parindent}{0pt}',r'\setlength{\parskip}{6pt}',r'\setlength{\emergencystretch}{3em}',r'\renewcommand{\arraystretch}{1.2}',r'\begin{document}',r'\begin{center}{\Large\bfseries '+esc(data['title'])+r'}\par\vspace{8pt}'+esc(data['author'])+r'\par '+esc(data['date'])+r'\end{center}']
    tex.insert(4,r'\DeclareUnicodeCharacter{2212}{\ensuremath{-}}')
    tex.insert(5,r'\DeclareUnicodeCharacter{03B2}{\ensuremath{\beta}}')
    tex.insert(6,r'\DeclareUnicodeCharacter{2265}{\ensuremath{\geq}}')
    tex.insert(7,r'\DeclareUnicodeCharacter{2264}{\ensuremath{\leq}}')
    tex.insert(8,r'\DeclareUnicodeCharacter{2248}{\ensuremath{\approx}}')
    tex.insert(9,r'\DeclareUnicodeCharacter{0394}{\ensuremath{\Delta}}')
    def render_figure(block):
        doc.add_page_break()
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(M/'figures'/block['file']),width=Inches(6.65))
        p=doc.add_paragraph(block['caption']);p.paragraph_format.space_after=Pt(10)
        for r in p.runs:r.font.size=Pt(9)
        vector=Path(block['file']).with_suffix('.pdf')
        figure_file=str(vector) if (M/'figures'/vector).exists() else block['file']
        tex.extend([r'\clearpage',r'\begin{figure}[p]\centering',r'\includegraphics[width=\linewidth]{figures/'+figure_file+'}',r'\par\small '+esc(block['caption']),r'\end{figure}',r'\clearpage'])
    eni=0
    figures=[block for block in data['sections'] if block['kind']=='figure']
    for block in [block for block in data['sections'] if block['kind']!='figure']:
        kind=block['kind']
        if kind=='section':
            doc.add_heading(block['title'],level=1);tex.append(r'\section*{'+esc(block['title'])+'}')
            for text in block['paragraphs']:
                p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;tex.append(esc(text)+'\n')
        elif kind=='table':
            p=doc.add_paragraph(block['caption']);p.paragraph_format.keep_with_next=True;p.runs[0].bold=True;p.runs[0].font.size=Pt(9)
            table=doc.add_table(rows=1,cols=len(block['headers']));table.autofit=False
            width=6.8/len(block['headers'])
            for cell,text in zip(table.rows[0].cells,block['headers']):cell.text=text
            for row in block['rows']:
                for cell,text in zip(table.add_row().cells,row):cell.text=str(text)
            for i,row in enumerate(table.rows):
                no_split(row)
                for cell in row.cells:
                    cell.width=Inches(width)
                    tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcPr.append(borders);margins=OxmlElement('w:tcMar')
                    for edge in ['top','bottom','left','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                    tcPr.append(margins)
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');tcPr.append(sh)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1
                        p.paragraph_format.keep_with_next=i<len(table.rows)-1
                        for r in p.runs:r.font.size=Pt(8.5);r.bold=i==0
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            tex.extend([r'\begin{table}[H]\small',r'\textbf{'+esc(block['caption'])+r'}\par\vspace{5pt}',r'\begin{tabularx}{\linewidth}{'+(r'>{\raggedright\arraybackslash}X'*len(block['headers']))+'}',r'\toprule',' & '.join(r'\textbf{'+esc(h)+'}' for h in block['headers'])+r'\\\midrule'])
            tex.extend(' & '.join(esc(x) for x in row)+r'\\' for row in block['rows']);tex.extend([r'\bottomrule\end{tabularx}',r'\end{table}'])
        elif kind=='equation':
            eni+=1;eqpath=M/'figures'/f'equation{eni}.png'
            fig=plt.figure(figsize=(7.6,.65));fig.text(.5,.5,'$'+block['latex']+'$',ha='center',va='center',fontsize=12);fig.savefig(eqpath,dpi=240,bbox_inches='tight',pad_inches=.1);plt.close(fig)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(eqpath),width=Inches(min(6.55,Image.open(eqpath).width/240)))
            tex.append(r'\['+block['latex']+r'\]')
    doc.add_heading('References',level=1);tex.append(r'\section*{References}\small')
    for ref in data['references']:
        p=doc.add_paragraph('['+ref['id']+'] '+ref['text']+' '+ref['url'])
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(9)
        tex.append(r'\begin{samepage}['+ref['id']+'] '+esc(ref['text'])+r' \url{'+ref['url']+r'}\par\end{samepage}')
    for block in figures:render_figure(block)
    tex.append(r'\end{document}');(M/'paper.tex').write_text('\n'.join(tex),encoding='utf8');doc.save(M/'paper.docx')
    print('Generated paper.docx and paper.tex from paper_content.json')

def render_external(output_dir):
    """Copy only small published presentation inputs; render outside the repository."""
    import shutil
    from data_paths import resolve_data_root
    global M
    source=M
    destination=resolve_data_root(output_dir,for_write=True)
    if destination.exists() and any(destination.iterdir()):
        raise FileExistsError('Manuscript output directory must be empty')
    destination.mkdir(parents=True,exist_ok=True)
    (destination/'figures').mkdir(exist_ok=True)
    data=json.loads((source/'paper_content.json').read_text(encoding='utf-8'))
    shutil.copyfile(source/'paper_content.json',destination/'paper_content.json')
    for block in data['sections']:
        if block.get('kind')!='figure':continue
        filename=Path(block['file'])
        if filename.is_absolute() or filename.name!=str(filename):raise ValueError('Figure filename must stay local')
        for name in [filename,filename.with_suffix('.pdf')]:
            original=source/'figures'/name
            if original.is_file():shutil.copyfile(original,destination/'figures'/name)
    M=destination
    try:main()
    finally:M=source


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True)
    args=parser.parse_args();render_external(args.output_dir)
